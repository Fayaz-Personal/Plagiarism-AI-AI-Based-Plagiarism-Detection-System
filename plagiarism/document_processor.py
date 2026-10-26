import os
import uuid
import hashlib
from pathlib import Path
from werkzeug.utils import secure_filename
import pypdf
import docx


class DocumentExtractionError(Exception):
    """Exception raised when document text extraction fails."""
    pass


def compute_file_hash(file_path: str) -> str:
    """Computes a SHA-256 cryptographic hash of a file for duplicate detection."""
    sha256 = hashlib.sha256()
    with open(file_path, 'rb') as f:
        while chunk := f.read(65536):
            sha256.update(chunk)
    return sha256.hexdigest()


def extract_text_from_pdf(file_path: str) -> tuple[str, bool, int]:
    """
    Extract text and page metadata from a PDF file using pypdf.
    Returns: (extracted_text, is_scanned_or_image_pdf, page_count)
    """
    path = Path(file_path)
    if not path.is_file():
        raise DocumentExtractionError(f"PDF file not found: {file_path}")

    extracted_pages = []
    page_count = 0

    try:
        reader = pypdf.PdfReader(str(path))
        page_count = len(reader.pages)

        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise DocumentExtractionError("The PDF document is password-protected and cannot be analyzed.")

        for idx, page in enumerate(reader.pages):
            try:
                page_text = page.extract_text() or ""
                if page_text.strip():
                    extracted_pages.append(page_text.strip())
            except Exception:
                continue

    except Exception as e:
        if isinstance(e, DocumentExtractionError):
            raise
        raise DocumentExtractionError(f"Failed to read PDF document: {str(e)}")

    full_text = "\n\n".join(extracted_pages).strip()

    # Check if PDF contains no extractable text (scanned image document)
    if not full_text:
        return "", True, max(1, page_count)

    return full_text, False, max(1, page_count)


def extract_text_from_docx(file_path: str) -> tuple[str, bool, int]:
    """
    Extract text from a Microsoft Word (.docx) file.
    Extracts body paragraphs and table contents.
    Returns: (extracted_text, is_empty, page_count)
    """
    path = Path(file_path)
    if not path.is_file():
        raise DocumentExtractionError(f"Word document not found: {file_path}")

    try:
        doc = docx.Document(str(path))
        content_parts = []

        # Extract paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                content_parts.append(text)

        # Extract table contents
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    content_parts.append(" | ".join(row_cells))

        full_text = "\n\n".join(content_parts).strip()
        is_empty = len(full_text) == 0
        return full_text, is_empty, 1

    except Exception as e:
        raise DocumentExtractionError(f"Failed to read Word (.docx) document: {str(e)}")


def extract_text_from_txt(file_path: str) -> tuple[str, bool, int]:
    """
    Extract text from a plain text file, handling multiple encodings and UTF-8 BOM.
    Normalizes line endings to LF.
    Returns: (extracted_text, is_empty, page_count)
    """
    path = Path(file_path)
    if not path.is_file():
        raise DocumentExtractionError(f"Text file not found: {file_path}")

    encodings = ['utf-8-sig', 'utf-8', 'latin-1', 'cp1252', 'iso-8859-1']
    content = None

    for enc in encodings:
        try:
            with open(path, 'r', encoding=enc) as f:
                content = f.read()
                break
        except (UnicodeDecodeError, UnicodeError):
            continue

    if content is None:
        # Fallback with error replacement
        with open(path, 'r', encoding='utf-8', errors='replace') as f:
            content = f.read()

    # Normalize line endings
    content = content.replace('\r\n', '\n').replace('\r', '\n').strip()
    is_empty = len(content) == 0
    return content, is_empty, 1


def extract_text(file_path: str, file_type: str) -> tuple[str, str | None, bool, int]:
    """
    Dispatcher to extract text based on file type.
    Returns: (extracted_text, error_message, is_scanned_notice, page_count)
    """
    ftype = file_type.lower().strip().lstrip('.')
    try:
        if ftype == 'pdf':
            text, is_scanned, pages = extract_text_from_pdf(file_path)
            if is_scanned:
                return "", "This PDF contains no extractable text layer. Scanned PDFs require OCR preprocessing before textual analysis.", True, pages
            return text, None, False, pages
        elif ftype == 'docx':
            text, is_empty, pages = extract_text_from_docx(file_path)
            if is_empty:
                return "", "The Word (.docx) document contains no extractable text.", False, 1
            return text, None, False, pages
        elif ftype == 'txt':
            text, is_empty, pages = extract_text_from_txt(file_path)
            if is_empty:
                return "", "The text document is empty.", False, 1
            return text, None, False, pages
        else:
            return "", f"Unsupported file extension: .{ftype}", False, 1
    except DocumentExtractionError as e:
        return "", str(e), False, 1
    except Exception as e:
        return "", f"Unexpected extraction error: {str(e)}", False, 1


def generate_safe_filename(original_filename: str) -> tuple[str, str]:
    """
    Generates a secure server-side stored filename using a UUID prefix and sanitized name.
    Returns: (stored_filename, file_extension)
    """
    safe_base = secure_filename(original_filename)
    if not safe_base:
        safe_base = "uploaded_document"

    ext = safe_base.rsplit('.', 1)[-1].lower() if '.' in safe_base else ""
    unique_prefix = uuid.uuid4().hex[:12]
    stored_filename = f"{unique_prefix}_{safe_base}"
    return stored_filename, ext


def is_allowed_file(filename: str, allowed_extensions: set[str]) -> bool:
    """Checks if the file extension is allowed."""
    if not filename or '.' not in filename:
        return False
    ext = filename.rsplit('.', 1)[-1].lower()
    return ext in allowed_extensions
