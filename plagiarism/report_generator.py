import os
from datetime import datetime
from xml.sax.saxutils import escape
from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, PageBreak, HRFlowable
)
from reportlab.pdfgen import canvas
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from models.document_models import Analysis, SimilarityResult, SentenceMatch


class NumberedCanvas(canvas.Canvas):
    """Two-pass canvas for dynamic 'Page X of Y' footers and headers."""
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748b"))

        # Top running header
        self.drawString(36, 756, "PlagiarismAI | Academic Document Verification Platform")
        self.setStrokeColor(colors.HexColor("#e2e8f0"))
        self.setLineWidth(0.5)
        self.line(36, 750, 576, 750)

        # Bottom footer
        self.line(36, 45, 576, 45)
        self.drawString(36, 32, "Confidential Academic Verification Report - Generated for Originality Assessment")
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(576, 32, page_str)
        self.restoreState()


def safe_text(text: str) -> str:
    """Escapes HTML entities and replaces characters that cannot be encoded in standard Latin-1."""
    if not text:
        return ""
    # Convert exotic Unicode glyphs safely to prevent ReportLab encoding crashes
    cleaned = str(text).encode('latin-1', 'replace').decode('latin-1')
    return escape(cleaned)


def generate_pdf_report(analysis: Analysis, output_path: str) -> str:
    """
    Generates a professional, publication-quality academic plagiarism report PDF using ReportLab.
    Includes repeated table headers, 6 executive metrics, passage classifications, and methodology details.
    """
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()

    # Custom typography styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#0f172a')
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#64748b')
    )
    heading2_style = ParagraphStyle(
        'Heading2_Custom',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12.5,
        leading=16,
        textColor=colors.HexColor('#1e293b'),
        spaceBefore=14,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        'Body_Custom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#334155')
    )
    body_bold = ParagraphStyle(
        'Body_Bold',
        parent=body_style,
        fontName='Helvetica-Bold'
    )
    table_cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor('#1e293b')
    )
    table_cell_header = ParagraphStyle(
        'TableCellHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=12,
        textColor=colors.white
    )
    disclaimer_style = ParagraphStyle(
        'Disclaimer',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=8,
        leading=11,
        textColor=colors.HexColor('#475569')
    )

    story = []

    # Title & Banner
    story.append(Paragraph("PlagiarismAI Analysis Report", title_style))
    story.append(Paragraph("Verify Originality. Protect Academic Integrity. | Comprehensive Similarity Breakdown", subtitle_style))
    story.append(Spacer(1, 12))

    # Analysis Overview Table
    doc_name = analysis.document.original_filename if analysis.document else "Unknown Document"
    analysis_date_str = analysis.analysis_date.strftime("%Y-%m-%d %H:%M:%S UTC")

    meta_data = [
        [
            Paragraph("<b>Document Analyzed:</b>", body_style),
            Paragraph(safe_text(doc_name), body_style),
            Paragraph("<b>Analysis ID:</b>", body_style),
            Paragraph(f"#{analysis.id}", body_style)
        ],
        [
            Paragraph("<b>Analysis Date:</b>", body_style),
            Paragraph(analysis_date_str, body_style),
            Paragraph("<b>Status:</b>", body_style),
            Paragraph(f"<font color='#059669'><b>{safe_text(analysis.analysis_status.upper())}</b></font>", body_style)
        ]
    ]

    meta_table = Table(meta_data, colWidths=[110, 180, 90, 160])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f8fafc')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#e2e8f0')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 12))

    # Methodology Summary Box
    methodology_text = (
        "<b>Detection Methodology:</b> Document and passage similarities are evaluated using "
        "scikit-learn TF-IDF vectorization (unigrams and bigrams, sublinear TF scaling) and Cosine Similarity "
        "across a shared vector feature space. Sentence pairs with cosine similarity exceeding threshold (&ge; 0.70) "
        "are aligned and extracted."
    )
    story.append(Paragraph(methodology_text, body_style))
    story.append(Spacer(1, 10))

    # Summary Statistics Table (6 Core Metrics)
    story.append(Paragraph("Executive Summary Metrics", heading2_style))

    highest_score = analysis.highest_similarity_score
    highest_color = '#10b981'
    if highest_score >= 60:
        highest_color = '#ef4444'
    elif highest_score >= 30:
        highest_color = '#f59e0b'
    elif highest_score >= 10:
        highest_color = '#3b82f6'

    orig_score = getattr(analysis, 'originality_score', round(max(0.0, 100.0 - highest_score), 1))

    summary_data = [
        [
            Paragraph("Max Similarity", body_bold),
            Paragraph("Avg Similarity", body_bold),
            Paragraph("References", body_bold),
            Paragraph("Matched Sents", body_bold),
            Paragraph("Text Coverage", body_bold),
            Paragraph("Originality Index", body_bold)
        ],
        [
            Paragraph(f"<font color='{highest_color}' size='12'><b>{highest_score:.1f}%</b></font><br/>({analysis.similarity_category})", table_cell_style),
            Paragraph(f"<font size='12'><b>{analysis.average_similarity_score:.1f}%</b></font>", table_cell_style),
            Paragraph(f"<font size='12'><b>{analysis.references_compared}</b></font>", table_cell_style),
            Paragraph(f"<font size='12'><b>{analysis.matching_sentence_count}</b></font>", table_cell_style),
            Paragraph(f"<font size='12'><b>{analysis.estimated_matching_text_coverage:.1f}%</b></font>", table_cell_style),
            Paragraph(f"<font color='#059669' size='12'><b>{orig_score:.1f}%</b></font>", table_cell_style)
        ]
    ]

    summary_table = Table(summary_data, colWidths=[95, 85, 80, 90, 95, 95])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0f172a')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor('#ffffff')),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#cbd5e1')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 12))

    # Reference Comparison Table
    story.append(Paragraph("Reference Document Comparisons", heading2_style))

    ref_headers = [
        Paragraph("Reference Document", table_cell_header),
        Paragraph("Document Similarity", table_cell_header),
        Paragraph("Matching Sentences", table_cell_header),
        Paragraph("Originality Indicator", table_cell_header)
    ]
    ref_rows = [ref_headers]

    results = SimilarityResult.query.filter_by(analysis_id=analysis.id).order_by(
        SimilarityResult.document_similarity_score.desc()
    ).all()

    if not results:
        ref_rows.append([
            Paragraph("No reference comparisons available.", table_cell_style),
            Paragraph("N/A", table_cell_style),
            Paragraph("0", table_cell_style),
            Paragraph("N/A", table_cell_style)
        ])
    else:
        for r in results:
            orig_index = round(100.0 - r.document_similarity_score, 1)
            ref_rows.append([
                Paragraph(safe_text(r.reference_document_name), table_cell_style),
                Paragraph(f"<b>{r.document_similarity_score:.1f}%</b>", table_cell_style),
                Paragraph(str(r.matching_sentence_count), table_cell_style),
                Paragraph(f"{orig_index:.1f}%", table_cell_style)
            ])

    ref_table = Table(ref_rows, colWidths=[240, 100, 100, 100], repeatRows=1)
    ref_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e293b')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#cbd5e1')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(ref_table)
    story.append(Spacer(1, 14))

    # Matching Passages Section
    story.append(Paragraph("Flagged Matching Passages", heading2_style))

    matches = SentenceMatch.query.filter_by(analysis_id=analysis.id).order_by(
        SentenceMatch.similarity_score.desc()
    ).all()

    if not matches:
        story.append(Paragraph("<i>No sentences exceeded the matching threshold against reference documents.</i>", body_style))
    else:
        passage_headers = [
            Paragraph("#", table_cell_header),
            Paragraph("Submitted Passage vs. Reference Source", table_cell_header),
            Paragraph("Similarity & Category", table_cell_header)
        ]
        passage_rows = [passage_headers]

        for idx, m in enumerate(matches, start=1):
            pct = m.percentage_score
            cat = getattr(m, 'match_category', 'Moderate Match')
            expl = getattr(m, 'match_explanation', 'Textual overlap detected.')

            sub_content = f"<b>Submitted:</b> {safe_text(m.submitted_sentence)}"
            ref_content = f"<b>Source [{safe_text(m.reference_document_name)}]:</b> {safe_text(m.reference_sentence)}"
            meta_info = f"<i>Classification:</i> {safe_text(cat)} &bull; <i>Note:</i> {safe_text(expl)}"
            combined = f"{sub_content}<br/><br/>{ref_content}<br/><br/>{meta_info}"

            passage_rows.append([
                Paragraph(str(idx), table_cell_style),
                Paragraph(combined, table_cell_style),
                Paragraph(f"<b>{pct:.1f}%</b><br/>{safe_text(cat)}", table_cell_style)
            ])

        match_table = Table(passage_rows, colWidths=[25, 415, 100], repeatRows=1)
        match_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#334155')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#cbd5e1')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ]))
        story.append(match_table)

    story.append(Spacer(1, 16))

    # Academic Disclaimer & Limitations Table
    disclaimer_box = [
        [
            Paragraph("<b>Academic Integrity Disclaimer & Limitations:</b>", body_bold)
        ],
        [
            Paragraph(
                "Similarity scores indicate textual resemblance computed via TF-IDF vectorization and cosine similarity, "
                "not definitive proof of academic plagiarism. Legitimate textual overlap can occur through standard domain terminology, "
                "properly cited scholarly quotations, common methodological descriptions, and bibliography formatting. "
                "All flagged passages must be manually evaluated by an academic reviewer to differentiate between legitimate citation "
                "and unauthorized copying.",
                disclaimer_style
            )
        ]
    ]
    disclaimer_table = Table(disclaimer_box, colWidths=[540])
    disclaimer_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f1f5f9')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#cbd5e1')),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(KeepTogether([disclaimer_table]))

    # Build the PDF document with two-pass NumberedCanvas
    doc.build(story, canvasmaker=NumberedCanvas)
    return output_path


def generate_docx_report(analysis: Analysis, output_path: str) -> str:
    """
    Generates a structured Microsoft Word (.docx) report using python-docx.
    """
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc = docx.Document()

    # Title & Subtitle
    title = doc.add_heading("PlagiarismAI Analysis Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    sub = doc.add_paragraph("Verify Originality. Protect Academic Integrity. - Comprehensive Similarity Assessment")
    sub.italic = True

    # Document details
    doc.add_heading("1. Analysis Overview", level=1)
    p = doc.add_paragraph()
    doc_name = analysis.document.original_filename if analysis.document else "Unknown Document"
    p.add_run(f"Document Name: {doc_name}\n").bold = True
    p.add_run(f"Analysis ID: #{analysis.id}\n")
    p.add_run(f"Analysis Date: {analysis.analysis_date.strftime('%Y-%m-%d %H:%M:%S UTC')}\n")
    p.add_run(f"Processing Status: {analysis.analysis_status.upper()}\n")

    # Metrics Summary
    doc.add_heading("2. Executive Summary Metrics", level=1)
    table = doc.add_table(rows=2, cols=6)
    table.style = 'Light Grid Accent 1'
    headers = ["Max Similarity", "Avg Similarity", "References", "Matched Sents", "Text Coverage", "Originality Index"]
    for i, h in enumerate(headers):
        table.cell(0, i).text = h

    orig_val = getattr(analysis, 'originality_score', round(max(0.0, 100.0 - analysis.highest_similarity_score), 1))
    vals = [
        f"{analysis.highest_similarity_score:.1f}% ({analysis.similarity_category})",
        f"{analysis.average_similarity_score:.1f}%",
        str(analysis.references_compared),
        str(analysis.matching_sentence_count),
        f"{analysis.estimated_matching_text_coverage:.1f}%",
        f"{orig_val:.1f}%"
    ]
    for i, v in enumerate(vals):
        table.cell(1, i).text = v

    # Reference documents
    doc.add_heading("3. Reference Comparisons", level=1)
    results = SimilarityResult.query.filter_by(analysis_id=analysis.id).all()
    ref_table = doc.add_table(rows=1, cols=3)
    ref_table.style = 'Table Grid'
    ref_headers = ["Reference Document", "Document Similarity", "Matching Sentences"]
    for i, h in enumerate(ref_headers):
        ref_table.cell(0, i).text = h

    for r in results:
        row = ref_table.add_row()
        row.cells[0].text = r.reference_document_name
        row.cells[1].text = f"{r.document_similarity_score:.1f}%"
        row.cells[2].text = str(r.matching_sentence_count)

    # Matched Sentences
    doc.add_heading("4. Flagged Matching Passages", level=1)
    matches = SentenceMatch.query.filter_by(analysis_id=analysis.id).all()
    if not matches:
        doc.add_paragraph("No matching sentences found exceeding threshold.")
    else:
        for idx, m in enumerate(matches, start=1):
            mp = doc.add_paragraph()
            cat = getattr(m, 'match_category', 'Moderate Match')
            mp.add_run(f"Match #{idx} - Similarity: {m.percentage_score:.1f}% ({cat})\n").bold = True
            mp.add_run(f"Submitted Passage: \"{m.submitted_sentence}\"\n")
            mp.add_run(f"Reference [{m.reference_document_name}]: \"{m.reference_sentence}\"\n")

    # Disclaimer
    doc.add_heading("5. Academic Evaluation Disclaimer", level=2)
    disc = doc.add_paragraph(
        "Similarity scores indicate textual resemblance, not proof of academic misconduct. "
        "Standard scholarly quotes, domain terminology, and citation formatting may influence results. "
        "All flagged passages should be manually evaluated by an academic reviewer."
    )
    disc.italic = True

    doc.save(output_path)
    return output_path
